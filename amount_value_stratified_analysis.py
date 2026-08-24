"""同额度分层验证（汇报文档第六节）：价值信号在额度之外。

分层口径：
- 平均收入、可支配盈余与样本数交叉表：按申请金额（requested_loan_amount）分层，
  控制的是客户自己申请的额度量级；
- 3M30 违约风险、可支配盈余交叉表：按总额度（total_amount）分层，
  控制的是实际授信额度（与"模型是否在学习历史额度策略"的质疑对应）。

可支配盈余（net_surplus）：收入扣除支出后的可支配盈余，全样本 100% 覆盖
（含未放款申请），故总额度口径的盈余表可覆盖全部 6 个额度区间，
弥补 3M30 在高额度段无有效观察的缺口。

分段：证据核心区按固定 500 步长切 500-2500；稀疏尾部合并为 2500-5000、5000+
（独立产品档）两段；0-500 无样本（申请金额下限 500）。

用法：python amount_value_stratified_analysis.py

输出（写入 output/model_analysis_20260429/）：
- amount_value_stratified_pivot.xlsx
    Sheet「样本数-申请金额 / 平均收入-申请金额 / 平均可支配盈余-申请金额 /
    平均可支配盈余-成交样本 / 3M30逾期率-总额度 / 3M30有效样本数-总额度 /
    平均可支配盈余-总额度」
- amount_value_income_by_bin.png     平均收入多线图（X=申请金额区间）
- amount_value_surplus_by_bin.png    平均可支配盈余多线图（X=申请金额区间）
- amount_value_surplus_ta_by_bin.png 平均可支配盈余多线图（X=总额度区间）
- amount_value_duedate30_by_bin.png  3M30逾期率多线图（X=总额度区间）

说明：3M30 逾期标志仅对放款后有 3 个月表现期的申请有值（全样本中仅约 1.3 万条，
集中在总额度 2500 以下），高额度段逾期率无法计算属数据现实，须在报告中披露。
"""

from __future__ import annotations

import logging
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from config import CONFIG
from model_bin_analysis import (
    build_analysis_dataset,
    calculate_group_metrics,
    resolve_config_paths,
    setup_logger,
)

PRIMARY_BIN = "primary_model_score_manual_bin"
AMOUNT_FIELD = "requested_loan_amount"
RISK_AMOUNT_FIELD = "total_amount"
AMOUNT_BUCKET = "amount_bucket"
RISK_AMOUNT_BUCKET = "risk_amount_bucket"
# 证据核心区（3M30 有效样本所在）按固定 500 步长：500-2500；
# 稀疏尾部合并：2500-5000 一段、5000+（独立产品档）一段；0-500 无样本不纳入。
AMOUNT_BUCKETS = [
    (500, 1000),
    (1000, 1500),
    (1500, 2000),
    (2000, 2500),
    (2500, 5000),
    (5000, None),
]
BIN_LABELS = [1, 2, 3, 4, 5]

CNT_METRIC = "sample_cnt"
INCOME_METRIC = "avg_total_income"
SURPLUS_METRIC = "avg_net_surplus"
DEAL_FLAG = "is_deal_application"
RISK_METRIC = "duedate_3m_30_bad_rate"
RISK_VALID_METRIC = "duedate_3m_30_valid_cnt"

OUTPUT_XLSX = "amount_value_stratified_pivot.xlsx"
OUTPUT_INCOME_PNG = "amount_value_income_by_bin.png"
OUTPUT_SURPLUS_PNG = "amount_value_surplus_by_bin.png"
OUTPUT_SURPLUS_TA_PNG = "amount_value_surplus_ta_by_bin.png"
OUTPUT_RISK_PNG = "amount_value_duedate30_by_bin.png"


def build_amount_bucket(amount: pd.Series) -> pd.Series:
    bins = [lo for lo, _ in AMOUNT_BUCKETS]
    bins.append(np.inf)
    labels = [f"{lo}-{hi}" if hi else f"{lo}+" for lo, hi in AMOUNT_BUCKETS]
    return pd.cut(amount, bins=bins, labels=labels, right=False, include_lowest=True)


def build_pivot(
    long: pd.DataFrame,
    row_total: pd.DataFrame,
    col_total: pd.DataFrame,
    grand: pd.DataFrame,
    metric: str,
    bucket_field: str,
) -> pd.DataFrame:
    body = long.pivot(index=bucket_field, columns=PRIMARY_BIN, values=metric)
    body["合计"] = row_total.set_index(bucket_field)[metric].reindex(body.index)
    total_row = col_total.set_index(PRIMARY_BIN)[metric].reindex(BIN_LABELS)
    total_row["合计"] = grand.iloc[0][metric]
    return pd.concat([body, pd.DataFrame([total_row], index=["合计"])])


def format_workbook(path: Path, number_formats: dict[str, str]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for sheet_name, fmt in number_formats.items():
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, max_col=ws.max_column):
            for cell in row:
                if cell.value is not None:
                    cell.number_format = fmt
    wb.save(path)


def plot_metric(df: pd.DataFrame, title: str, ylabel: str, xlabel: str, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 5.5))
    x = np.arange(len(df.index))
    for label in BIN_LABELS:
        ax.plot(x, df[label].to_numpy(), marker="o", label=f"价值{label}")
    ax.set_xticks(x, df.index)
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.legend(title="价值分箱")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def main() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False

    cfg = resolve_config_paths(CONFIG)
    logger = setup_logger(level=cfg["project"].get("log_level", "INFO"))
    out_dir = Path(cfg["project"]["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)

    logger.info("building analysis dataset")
    df = build_analysis_dataset(cfg, logger)
    df[AMOUNT_BUCKET] = build_amount_bucket(pd.to_numeric(df[AMOUNT_FIELD], errors="coerce"))
    df[RISK_AMOUNT_BUCKET] = build_amount_bucket(pd.to_numeric(df[RISK_AMOUNT_FIELD], errors="coerce"))

    special_bins = cfg["score_binning"]["schemes"][0]["scores"][0].get("null_values", [-1])
    special_cnt = int(df[PRIMARY_BIN].isin(special_bins).sum())
    df = df[~df[PRIMARY_BIN].isin(special_bins)].copy()
    if special_cnt:
        logger.info(f"excluded {special_cnt:,} rows with special value bin {special_bins}")

    df = df[df[AMOUNT_BUCKET].notna() & df[RISK_AMOUNT_BUCKET].notna()].copy()

    long_income = calculate_group_metrics(df, [AMOUNT_BUCKET, PRIMARY_BIN], cfg, ["sample", "mean"])
    row_income = calculate_group_metrics(df, [AMOUNT_BUCKET], cfg, ["sample", "mean"])
    col_income = calculate_group_metrics(df, [PRIMARY_BIN], cfg, ["sample", "mean"])
    grand_income = calculate_group_metrics(df, [], cfg, ["sample", "mean"])

    long_risk = calculate_group_metrics(df, [RISK_AMOUNT_BUCKET, PRIMARY_BIN], cfg, ["risk"])
    row_risk = calculate_group_metrics(df, [RISK_AMOUNT_BUCKET], cfg, ["risk"])
    col_risk = calculate_group_metrics(df, [PRIMARY_BIN], cfg, ["risk"])
    grand_risk = calculate_group_metrics(df, [], cfg, ["risk"])

    # 可支配盈余按总额度分层的均值组（net_surplus 全样本 100% 覆盖，6 个区间均可计算）
    long_surplus_ta = calculate_group_metrics(df, [RISK_AMOUNT_BUCKET, PRIMARY_BIN], cfg, ["mean"])
    row_surplus_ta = calculate_group_metrics(df, [RISK_AMOUNT_BUCKET], cfg, ["mean"])
    col_surplus_ta = calculate_group_metrics(df, [PRIMARY_BIN], cfg, ["mean"])
    grand_surplus_ta = calculate_group_metrics(df, [], cfg, ["mean"])

    # 成交（放款）子样本上的可支配盈余 × 申请金额交叉表（口径与表4 一致，样本替换为已放款申请）
    deal_df = df[df[DEAL_FLAG].astype(bool)]
    logger.info(f"deal sample: n={len(deal_df):,}")
    long_surplus_deal = calculate_group_metrics(deal_df, [AMOUNT_BUCKET, PRIMARY_BIN], cfg, ["mean"])
    row_surplus_deal = calculate_group_metrics(deal_df, [AMOUNT_BUCKET], cfg, ["mean"])
    col_surplus_deal = calculate_group_metrics(deal_df, [PRIMARY_BIN], cfg, ["mean"])
    grand_surplus_deal = calculate_group_metrics(deal_df, [], cfg, ["mean"])

    bucket_counts = row_income.set_index(AMOUNT_BUCKET)[CNT_METRIC]
    for bucket in bucket_counts.index:
        logger.info(f"requested amount bucket {bucket}: n={int(bucket_counts[bucket]):,}")
    logger.info(f"total rows: {len(df):,}")
    risk_valid = row_risk.set_index(RISK_AMOUNT_BUCKET)[RISK_VALID_METRIC]
    thin_buckets = [b for b in risk_valid.index if int(risk_valid[b]) < 100]
    if thin_buckets:
        logger.warning(
            "3M30 valid samples < 100 in total-amount buckets: "
            + ", ".join(str(b) for b in thin_buckets)
        )

    cnt_pivot = build_pivot(long_income, row_income, col_income, grand_income, CNT_METRIC, AMOUNT_BUCKET)
    income_pivot = build_pivot(long_income, row_income, col_income, grand_income, INCOME_METRIC, AMOUNT_BUCKET)
    surplus_pivot = build_pivot(long_income, row_income, col_income, grand_income, SURPLUS_METRIC, AMOUNT_BUCKET)
    surplus_ta_pivot = build_pivot(long_surplus_ta, row_surplus_ta, col_surplus_ta, grand_surplus_ta, SURPLUS_METRIC, RISK_AMOUNT_BUCKET)
    surplus_deal_pivot = build_pivot(long_surplus_deal, row_surplus_deal, col_surplus_deal, grand_surplus_deal, SURPLUS_METRIC, AMOUNT_BUCKET)
    risk_pivot = build_pivot(long_risk, row_risk, col_risk, grand_risk, RISK_METRIC, RISK_AMOUNT_BUCKET)
    risk_valid_pivot = build_pivot(long_risk, row_risk, col_risk, grand_risk, RISK_VALID_METRIC, RISK_AMOUNT_BUCKET)

    # 稳健性核对：均值受极端值影响，同时输出中位数（不进 xlsx，日志核对趋势是否一致）
    surplus_median = (
        df.groupby([AMOUNT_BUCKET, PRIMARY_BIN], observed=True)["net_surplus"]
        .median()
        .unstack(PRIMARY_BIN)
    )
    logger.info("net_surplus 中位数核对（申请金额口径）：")
    logger.info("\n" + surplus_median.round(1).to_string())

    xlsx_path = out_dir / OUTPUT_XLSX
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        cnt_pivot.to_excel(writer, sheet_name="样本数-申请金额")
        income_pivot.to_excel(writer, sheet_name="平均收入-申请金额")
        surplus_pivot.to_excel(writer, sheet_name="平均可支配盈余-申请金额")
        risk_pivot.to_excel(writer, sheet_name="3M30逾期率-总额度")
        risk_valid_pivot.to_excel(writer, sheet_name="3M30有效样本数-总额度")
        surplus_ta_pivot.to_excel(writer, sheet_name="平均可支配盈余-总额度")
        surplus_deal_pivot.to_excel(writer, sheet_name="平均可支配盈余-成交样本")
    format_workbook(
        xlsx_path,
        {
            "样本数-申请金额": "#,##0",
            "平均收入-申请金额": "#,##0.0",
            "平均可支配盈余-申请金额": "#,##0.0",
            "3M30逾期率-总额度": "0.00%",
            "3M30有效样本数-总额度": "#,##0",
            "平均可支配盈余-总额度": "#,##0.0",
            "平均可支配盈余-成交样本": "#,##0.0",
        },
    )
    logger.info(f"workbook written: {xlsx_path}")

    plot_metric(
        income_pivot.drop(index="合计", columns="合计"),
        "各申请金额区间内不同价值分箱的平均收入",
        "平均收入",
        "申请金额区间",
        out_dir / OUTPUT_INCOME_PNG,
    )
    plot_metric(
        surplus_pivot.drop(index="合计", columns="合计"),
        "各申请金额区间内不同价值分箱的平均可支配盈余",
        "平均可支配盈余(net_surplus)",
        "申请金额区间",
        out_dir / OUTPUT_SURPLUS_PNG,
    )
    plot_metric(
        surplus_ta_pivot.drop(index="合计", columns="合计"),
        "各总额度区间内不同价值分箱的平均可支配盈余",
        "平均可支配盈余(net_surplus)",
        "总额度区间",
        out_dir / OUTPUT_SURPLUS_TA_PNG,
    )
    plot_metric(
        risk_pivot.drop(index="合计", columns="合计"),
        "各总额度区间内不同价值分箱的3M30逾期率",
        "3M30逾期率",
        "总额度区间",
        out_dir / OUTPUT_RISK_PNG,
    )
    logger.info(f"plots written: {out_dir}")


if __name__ == "__main__":
    main()
